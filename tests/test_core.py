from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from uuid import uuid4
import json
import re
import pytest
from openpyxl import load_workbook
from auxiliaries.validation import parse_date, parse_money, validate_row, REPORT_PATTERN
from auxiliaries.reconciliation import plan_pairs, classify_pair
from auxiliaries.config import ROOT, load_config, resolve_path
from auxiliaries.audit import Audit
from auxiliaries.excel import create_masters, read_master, write_report


@pytest.mark.parametrize('value,expected', [('1,234.56','1234.56'),('-25.00','-25.00'),('($1,000.20)','-1000.20'),('0.00','0.00')])
def test_money_valid(value, expected):
    assert parse_money(value) == Decimal(expected)


@pytest.mark.parametrize('value', ['', '1.234,56', '12,34.56','NaN','--1.00','(-1.00)','1.001','10000000000000000.00'])
def test_money_invalid(value):
    with pytest.raises(ValueError): parse_money(value)


@pytest.mark.parametrize('value', ['2026-02-30','2025-02-29','', '01/02/2026'])
def test_date_invalid(value):
    with pytest.raises(ValueError): parse_date(value)


def row(index=1, web_id=None, amount='10.00'):
    raw=dict(account='800003', date_raw='2026-10-01', description='  Pago   DEMO ',amount_raw=amount,
             row_index=index,web_id=web_id,source_url='SINTETICO')
    return validate_row(raw,'00000000-0000-4000-8000-000000000001')['row']


RULES={'800003':dict(enabled=True,tolerance='0.00')}


def statuses(w,m,rules=RULES):
    return [classify_pair(p) for p in plan_pairs(w,m,rules)]


def test_classifications_and_duplicates():
    assert statuses([row()], [row()]) == ['CONCILIADO']
    assert statuses([row()], []) == ['SOLO_WEB']
    assert statuses([], [row()]) == ['SOLO_MASTER']
    assert statuses([row(web_id='A')],[row(web_id='A',amount='11.00')]) == ['DISCREPANCIA']
    assert statuses([row(),row(2)],[row(),row(2)]) == ['CONCILIADO','CONCILIADO']
    assert statuses([row(),row(2)],[row()]) == ['AMBIGUO','AMBIGUO']
    assert statuses([row(web_id='A'),row(2,web_id='A')],[row(web_id='A')]) == ['AMBIGUO','AMBIGUO']


def test_sql_rule_changes_classification():
    assert statuses([row(web_id='A')],[row(web_id='A',amount='10.01')]) == ['DISCREPANCIA']
    assert statuses([row(web_id='A')],[row(web_id='A',amount='10.01')],{'800003':dict(enabled=True,tolerance='0.01')}) == ['CONCILIADO']
    with pytest.raises(ValueError): plan_pairs([row()],[],{'800003':dict(enabled=False,tolerance='0.00')})


def config_tmp(tmp_path):
    cfg=load_config(require_secrets=False)
    cfg.update(PROJECT_ROOT=str(tmp_path),output_path=str(tmp_path/'output'))
    return cfg


def test_audit_error_outbox_redaction_and_ids(tmp_path):
    cfg=config_tmp(tmp_path); cfg['DB_PASSWORD']='private-password'
    a=Audit(cfg); b=Audit(cfg)
    assert a.id != b.id
    with pytest.raises(RuntimeError):
        with a.step('SQL'):
            raise RuntimeError('private-password failed')
    e=a.events[0]
    assert e['status']=='ERROR' and e['duration_ms']>=0 and e['ended_at']
    assert 'private-password' not in (a.folder/'audit.jsonl').read_text()
    assert len(a.pending())==1
    with pytest.raises(ValueError): a.finish('OK')
    a.finish('PARCIAL')


def test_master_and_report(tmp_path):
    create_masters(tmp_path)
    p=tmp_path/'data/conciliacion_master.xlsx'
    with pytest.raises(ValueError): read_master(p,str(uuid4()),{'800003'},'2026-10-01','2026-10-31')
    master=read_master(p,str(uuid4()),{'800003'},'2026-10-01','2026-10-31',True)
    assert len(master['rows'])==3 and len(master['metadata']['sha256'])==64
    with pytest.raises(ValueError): read_master(p,str(uuid4()),{'800011'},'2026-10-01','2026-10-31',True)
    malicious=row(); malicious['description']='=HYPERLINK("http://example.invalid")'
    dest=tmp_path/'Reporte_TestFire_20261016_101010_00000000-0000-4000-8000-000000000001.xlsx'
    assert re.fullmatch(REPORT_PATTERN,dest.name)
    write_report(dest,[malicious],[],[],[],{'status':'PARCIAL'})
    wb=load_workbook(dest)
    assert wb.sheetnames==['Movimientos','Conciliacion','Rechazados','Auditoria','Resumen']
    assert not any(c.data_type=='f' for ws in wb for rr in ws for c in rr)
    cols={c.value:c.column for c in wb['Movimientos'][1]}
    assert wb['Movimientos'].cell(2,cols['amount']).data_type=='n'
    assert wb['Movimientos'].cell(2,cols['date']).is_date
    wb.close()


def test_path_traversal(tmp_path):
    with pytest.raises(ValueError): resolve_path(tmp_path,'../escape')
    with pytest.raises(ValueError): resolve_path(tmp_path,str(tmp_path/'absolute'))
    assert resolve_path(tmp_path,'data/a.xlsx')==tmp_path/'data/a.xlsx'
