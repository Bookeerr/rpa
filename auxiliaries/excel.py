from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import hashlib
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from auxiliaries.validation import validate_row

COLUMNS = ["account", "date_raw", "description", "amount_raw", "web_id", "currency"]


def create_masters(root):
    folder = Path(root) / "data"
    folder.mkdir(parents=True, exist_ok=True)
    for name, synthetic in (("conciliacion_master.xlsx", True), ("plantilla_master_real.xlsx", False)):
        path = folder / name
        if path.exists():
            continue
        wb = Workbook()
        ws = wb.active
        ws.title = "Master"
        ws.append(COLUMNS)
        if synthetic:
            for row in [ ["800003", "2026-10-01", "Deposito demo", "100.00", "DEMO-1", None],
                         ["800003", "2026-10-02", "Cargo demo", "-25.00", "DEMO-2", None],
                         ["800003", "2026-10-03", "Solo maestro", "7.00", "DEMO-3", None] ]:
                ws.append(row)
        meta = wb.create_sheet("Metadata")
        for row in [("key", "value"), ("version", "SINTETICO-1" if synthetic else "COMPLETAR"),
                    ("synthetic", "true" if synthetic else "false"),
                    ("period_start", "2026-10-01"), ("period_end", "2026-10-31")]:
            meta.append(row)
        wb.save(path)


def read_master(path, execution_id, enabled_accounts, period_start, period_end, allow_synthetic=False):
    from auxiliaries.audit import utc_now
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        meta = dict(list(wb["Metadata"].values)[1:])
        if meta.get("synthetic") == "true" and not allow_synthetic:
            raise ValueError("Maestro sintetico: habilitar solo en pruebas")
        if not meta.get("version") or meta["version"] == "COMPLETAR":
            raise ValueError("Completar version del maestro")
        if (meta.get("period_start"), meta.get("period_end")) != (period_start, period_end):
            raise ValueError("Periodo de maestro no coincide")
        start, end = date.fromisoformat(period_start), date.fromisoformat(period_end)
        if start > end:
            raise ValueError("Periodo invertido")
        values = iter(wb["Master"].values)
        if list(next(values)) != COLUMNS:
            raise ValueError("Columnas de maestro incorrectas")
        rows = []
        for index, cells in enumerate(values, 2):
            if all(v is None for v in cells):
                continue
            raw = dict(zip(COLUMNS, cells))
            raw["account"] = str(raw["account"])
            if isinstance(raw["date_raw"], (date, datetime)):
                raw["date_raw"] = raw["date_raw"].strftime("%Y-%m-%d")
            if isinstance(raw["amount_raw"], (int, float, Decimal)):
                raw["amount_raw"] = str(Decimal(str(raw["amount_raw"])).quantize(Decimal("0.01")))
            raw.update(row_index=index, source_url=None)
            result = validate_row(raw, execution_id)
            if not result["valid"]:
                raise ValueError(f"Maestro fila {index}: {result['error']}")
            row = result["row"]
            if row["account"] not in enabled_accounts or not start <= date.fromisoformat(row["date"]) <= end:
                raise ValueError(f"Maestro fila {index}: cuenta o periodo fuera del alcance")
            rows.append(row)
        return {"rows": rows, "metadata": dict(meta, name=path.name,
                sha256=hashlib.sha256(path.read_bytes()).hexdigest(), read_at=utc_now())}
    finally:
        wb.close()


def write_report(path, movements, reconciliation, rejected, audit, summary):
    """Escribe valores tipados. Todas las cadenas son texto, nunca formulas."""
    wb = Workbook()
    wb.remove(wb.active)
    datasets = {"Movimientos": movements, "Conciliacion": reconciliation,
                "Rechazados": rejected, "Auditoria": audit,
                "Resumen": [{"key": k, "value": v} for k, v in summary.items()]}
    for name, rows in datasets.items():
        ws = wb.create_sheet(name)
        columns = list(dict.fromkeys(k for row in rows for k in row)) or ["sin_datos"]
        ws.append(columns)
        for row in rows:
            for col, key in enumerate(columns, 1):
                value = row.get(key)
                if isinstance(value, (list, dict)):
                    value = json.dumps(value, ensure_ascii=False, default=str)
                if key == "date" and isinstance(value, str):
                    value = date.fromisoformat(value)
                if key in ("amount", "balance") and value is not None:
                    value = Decimal(str(value))
                cell = ws.cell(ws.max_row + 1 if col == 1 else ws.max_row, col, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                if isinstance(value, Decimal):
                    cell.number_format = '#,##0.00;[Red]-#,##0.00'
                if isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="16324F")
            ws.column_dimensions[cell.column_letter].width = min(48, max(18, len(str(cell.value)) + 3))
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.xlsx")
    wb.save(temporary)
    temporary.replace(path)
    wb.close()
    return str(path)
