"""Verifica estructura y produce vistas para revision visual, sin simular Office."""
import json
from pathlib import Path
import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/evidencias/revision'


def contact(images,dest,columns=2,width=500):
    thumbs=[]
    for p in images:
        im=Image.open(p).convert('RGB');im.thumbnail((width,650))
        thumbs.append((p.name,im))
    rows=(len(thumbs)+columns-1)//columns
    sheet=Image.new('RGB',(columns*width,rows*680),'#e2e8f0');draw=ImageDraw.Draw(sheet)
    for n,(name,im) in enumerate(thumbs):
        x=(n%columns)*width;y=(n//columns)*680
        draw.text((x+10,y+5),name,fill='black');sheet.paste(im,(x,y+25))
    sheet.save(dest)


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    checks={}
    for filename,prefix in [('Informe_Tecnico_EP2.pdf','informe'),('Presentacion_EP2.pdf','presentacion')]:
        doc=pymupdf.open(ROOT/'docs/entrega'/filename)
        images=[];pages=[]
        for n,page in enumerate(doc,1):
            assert len(page.get_text().strip())>100,(filename,n,'pagina vacia o solo pie')
            blocks=page.get_text('dict')['blocks']
            spans=[s for b in blocks if 'lines' in b for line in b['lines'] for s in line['spans'] if s['text'].strip()]
            overflow=[s['text'] for s in spans if not (page.rect+pymupdf.Rect(-1,-1,1,1)).contains(pymupdf.Rect(s['bbox']))]
            assert not overflow,(filename,n,overflow)
            image=OUT/f'{prefix}_{n:02}.png'
            page.get_pixmap(matrix=pymupdf.Matrix(min(1,1200/page.rect.width),min(1,1200/page.rect.width))).save(image)
            images.append(image)
            pages.append(dict(page=n,width=round(page.rect.width,2),height=round(page.rect.height,2),
                              min_font=round(min(s['size'] for s in spans),2),text_inside_page=True))
        for start in range(0,len(images),4):
            contact(images[start:start+4],OUT/f'{prefix}_contacto_{start//4+1}.png')
        checks[filename]=pages
        for old in OUT.glob(prefix+'_[0-9][0-9].png'):
            if int(old.stem.rsplit('_',1)[1])>len(pages):old.unlink()
        doc.close()
    prs=Presentation(ROOT/'docs/entrega/Presentacion_EP2.pptx')
    for n,slide in enumerate(prs.slides,1):
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width+10000 and shape.top+shape.height<=prs.slide_height+10000,('PPT fuera de lienzo',n)
        assert slide.notes_slide.notes_text_frame.text.strip()
    checks['pptx']={'slides':len(prs.slides),'shapes_inside_canvas':True,'speaker_notes':True,
                    'powerpoint_render':'PENDIENTE; PDF paralelo revisable, no export de Office'}
    (ROOT/'docs/evidencias/revision_documentos.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
    print(json.dumps({k:len(v) if isinstance(v,list) else v for k,v in checks.items()},indent=2))


if __name__=='__main__':main()
