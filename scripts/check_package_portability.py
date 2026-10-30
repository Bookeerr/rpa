"""Descomprime el paquete en ruta con espacios y prueba auxiliares desde otro cwd."""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]


def main():
    archive=ROOT/'entrega/TestFire_EP2_20261020.zip'
    with tempfile.TemporaryDirectory(prefix='EP2 entrega con espacios ') as temp:
        target=Path(temp)
        with zipfile.ZipFile(archive) as z:
            for name in z.namelist():
                assert (target/name).resolve().is_relative_to(target.resolve())
            z.extractall(target)
        extracted=target/'TestFire_EP2'
        env=dict(os.environ,PYTHONPATH=str(extracted))
        run=subprocess.run([sys.executable,'-m','scripts.demo'],cwd=target,env=env,
                           capture_output=True,text=True,encoding='utf-8',timeout=90)
        assert run.returncode==0,run.stderr
        result=json.loads(run.stdout)
        report=extracted/result['report']
        assert report.is_file() and report.resolve().is_relative_to(extracted.resolve())
        wb=load_workbook(report,read_only=True)
        assert wb.sheetnames==['Movimientos','Conciliacion','Rechazados','Auditoria','Resumen']
        sheets=wb.sheetnames;wb.close()
        result.update(archive=archive.relative_to(ROOT).as_posix(),extracted_path_with_spaces=True,
                      different_working_directory=True,excel_sheets=sheets,scope='AUXILIARES; no acredita otra PC ni Studio')
    (ROOT/'docs/evidencias/portabilidad_paquete.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))


if __name__=='__main__':main()
